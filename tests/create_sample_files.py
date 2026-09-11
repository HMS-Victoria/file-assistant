"""生成不含私人信息的 Milestone 2 测试样例文件。"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook


SAMPLE_DIRECTORY = Path(__file__).parent / "sample_files"


def create_samples() -> None:
    """创建最小、稳定的 Word、Excel、PDF 测试样例。"""
    SAMPLE_DIRECTORY.mkdir(exist_ok=True)

    word = Document()
    word.add_heading("会议通知", level=1)
    word.add_paragraph("会议时间是 2026 年 9 月 10 日。")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = "200"
    word.save(SAMPLE_DIRECTORY / "会议通知.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报销表"
    sheet.append(["姓名", "金额"])
    sheet.append(["张三", 200])
    workbook.create_sheet("说明").append(["这是测试文件"])
    workbook.save(SAMPLE_DIRECTORY / "报销表.xlsx")

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "报销通知\n请于 2026 年 9 月 10 日提交报销材料。")
    pdf.save(SAMPLE_DIRECTORY / "报销通知.pdf")
    pdf.close()


if __name__ == "__main__":
    create_samples()
