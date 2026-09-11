"""Milestone 5 的受限 Word、Excel 副本修改服务。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.operations.schemas import (
    AppendColumn,
    AppendRow,
    AppendText,
    DeleteColumn,
    DeleteRow,
    DeleteText,
    OperationPlan,
    ReplaceCell,
    ReplaceText,
)


class ModificationError(RuntimeError):
    """受限修改的安全错误。"""

    user_message = "这次不能安全地修改文件，请检查文件后再试。"


@dataclass(frozen=True)
class ModificationResult:
    """写入新文件后返回的变更摘要。"""

    output_path: Path
    changes: list[str]


class ModificationService:
    """只执行白名单操作，并且从不允许覆盖原文件。"""

    def apply(self, source_path: Path, output_path: Path, plan: OperationPlan) -> ModificationResult:
        """将同类型白名单操作安全写入新的 `.docx` 或 `.xlsx` 文件。"""
        source = source_path.resolve()
        output = output_path.resolve()
        if not source.is_file():
            raise ModificationError("原文件不存在。")
        if source == output:
            raise ModificationError("不允许覆盖原文件。")
        if output.exists():
            raise ModificationError("新文件名已存在，为避免覆盖请换一个名称。")
        if not output.parent.is_dir():
            raise ModificationError("保存位置不存在。")
        if source.suffix.lower() != output.suffix.lower() or source.suffix.lower() not in {".docx", ".xlsx"}:
            raise ModificationError("目前只能将 Word 或 Excel 修改为同类型的新文件。")

        if source.suffix.lower() == ".docx":
            return self._apply_word(source, output, plan)
        return self._apply_excel(source, output, plan)

    def _apply_word(self, source: Path, output: Path, plan: OperationPlan) -> ModificationResult:
        allowed = (ReplaceText, AppendText, DeleteText)
        if any(not isinstance(operation, allowed) for operation in plan.operations):
            raise ModificationError("这份 Word 文件只能执行文字替换、增加或删除文字。")
        document = Document(source)
        changes: list[str] = []
        for operation in plan.operations:
            if isinstance(operation, ReplaceText):
                count = self._replace_in_paragraphs(self._word_paragraphs(document), operation.target, operation.replacement)
                changes.append(f"将“{operation.target}”替换为“{operation.replacement}”（{count} 处）")
            elif isinstance(operation, DeleteText):
                count = self._replace_in_paragraphs(self._word_paragraphs(document), operation.target, "")
                changes.append(f"删除“{operation.target}”（{count} 处）")
            else:
                document.add_paragraph(operation.text)
                changes.append("在文档末尾增加了一段文字")
        self._save_document_safely(document, output)
        return ModificationResult(output, changes)

    @staticmethod
    def _word_paragraphs(document: Document) -> list[Paragraph]:
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        return paragraphs

    @staticmethod
    def _replace_in_paragraphs(paragraphs: list[Paragraph], target: str, replacement: str) -> int:
        count = 0
        for paragraph in paragraphs:
            occurrences = paragraph.text.count(target)
            if occurrences:
                paragraph.text = paragraph.text.replace(target, replacement)
                count += occurrences
        return count

    def _apply_excel(self, source: Path, output: Path, plan: OperationPlan) -> ModificationResult:
        allowed = (ReplaceCell, AppendRow, DeleteRow, AppendColumn, DeleteColumn)
        if any(not isinstance(operation, allowed) for operation in plan.operations):
            raise ModificationError("这份 Excel 文件只能执行单元格、行或列的基础修改。")
        workbook = load_workbook(source)
        changes: list[str] = []
        try:
            for operation in plan.operations:
                sheet = self._worksheet(workbook, operation.sheet)
                if isinstance(operation, ReplaceCell):
                    sheet[operation.cell] = operation.value
                    changes.append(f"将“{sheet.title}”的 {operation.cell} 修改为“{operation.value}”")
                elif isinstance(operation, AppendRow):
                    sheet.append(operation.values)
                    changes.append(f"在“{sheet.title}”末尾增加一行")
                elif isinstance(operation, DeleteRow):
                    if operation.row > sheet.max_row:
                        raise ModificationError(f"“{sheet.title}”没有第 {operation.row} 行。")
                    sheet.delete_rows(operation.row, 1)
                    changes.append(f"删除“{sheet.title}”的第 {operation.row} 行")
                elif isinstance(operation, AppendColumn):
                    column = sheet.max_column + 1
                    for row, value in enumerate(operation.values, start=1):
                        sheet.cell(row=row, column=column, value=value)
                    changes.append(f"在“{sheet.title}”末尾增加一列")
                else:
                    if operation.column > sheet.max_column:
                        raise ModificationError(f"“{sheet.title}”没有第 {operation.column} 列。")
                    sheet.delete_cols(operation.column, 1)
                    changes.append(f"删除“{sheet.title}”的第 {operation.column} 列")
            self._save_workbook_safely(workbook, output)
        finally:
            workbook.close()
        return ModificationResult(output, changes)

    @staticmethod
    def _worksheet(workbook, sheet_name: str):
        if sheet_name not in workbook.sheetnames:
            raise ModificationError(f"没有找到“{sheet_name}”这个工作表。")
        return workbook[sheet_name]

    @staticmethod
    def _temporary_output_path(output: Path) -> Path:
        with NamedTemporaryFile(prefix="file_assistant_", suffix=output.suffix, dir=output.parent, delete=False) as temporary:
            return Path(temporary.name)

    def _save_document_safely(self, document: Document, output: Path) -> None:
        temporary = self._temporary_output_path(output)
        try:
            document.save(temporary)
            temporary.replace(output)
        finally:
            temporary.unlink(missing_ok=True)

    def _save_workbook_safely(self, workbook, output: Path) -> None:
        temporary = self._temporary_output_path(output)
        try:
            workbook.save(temporary)
            temporary.replace(output)
        finally:
            temporary.unlink(missing_ok=True)
