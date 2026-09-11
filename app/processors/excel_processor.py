"""Excel .xlsx 文件的只读处理器。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.processors.base_processor import FileProcessor, FileReadResult, WorksheetData


class ExcelProcessor(FileProcessor):
    """读取工作簿、工作表和单元格值，不载入宏也不保存文件。"""

    supported_suffixes = frozenset({".xlsx"})
    kind = "excel"

    def read(self, file_path: Path) -> FileReadResult:
        self._validate_file(file_path)
        workbook = load_workbook(file_path, read_only=True, data_only=False)
        try:
            worksheets = [
                WorksheetData(
                    name=worksheet.title,
                    rows=[list(row) for row in worksheet.iter_rows(values_only=True)],
                )
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()
        return FileReadResult(file_path=file_path, kind="excel", worksheets=worksheets)
