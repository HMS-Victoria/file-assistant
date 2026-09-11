"""Word .docx 文件的只读处理器。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from app.processors.base_processor import FileProcessor, FileReadResult, TableData


class WordProcessor(FileProcessor):
    """提取 Word 段落与表格中的文字。"""

    supported_suffixes = frozenset({".docx"})
    kind = "word"

    def read(self, file_path: Path) -> FileReadResult:
        self._validate_file(file_path)
        document = Document(file_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [
            TableData(rows=[[cell.text for cell in row.cells] for row in table.rows])
            for table in document.tables
        ]
        return FileReadResult(file_path=file_path, kind="word", paragraphs=paragraphs, tables=tables)
