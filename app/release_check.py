"""便携包内置的离线验收：仅向全新指定目录写入合成样例与结果。"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import platform
import sys


def run(directory: str) -> int:
    root = Path(directory).resolve()
    root.mkdir(parents=True, exist_ok=False)
    os.environ["LOCALAPPDATA"] = str(root / "独立用户数据")
    report = {"platform": platform.platform(), "frozen": bool(getattr(sys, "frozen", False)),
              "executable": sys.executable, "checks": [], "real_clean_machine": False,
              "real_ai_account": False}
    try:
        import pymupdf
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from PySide6.QtWidgets import QApplication
        from app.config.settings import DeepSeekSettings, settings_path
        from app.operations.schemas import validate_operation_plan
        from app.processors.factory import get_processor
        from app.services.safe_modification_service import SafeModificationService
        from app.ui.main_window import MainWindow
        from app.ui.settings_dialog import SettingsDialog

        application = QApplication.instance() or QApplication([])
        report["qt_platform"] = application.platformName()
        samples = root / "中文 空格样例"
        samples.mkdir()
        word = samples / "会议 通知.docx"
        document = Document()
        document.add_paragraph("测试姓名甲，金额 200。")
        document.save(word)
        excel = samples / "报销 表.xlsx"
        workbook = Workbook()
        workbook.active.title = "报销表"
        workbook.active.append(["姓名", "金额"])
        workbook.active.append(["测试姓名甲", 200])
        workbook.save(excel)
        workbook.close()
        pdf_path = samples / "通知 PDF.pdf"
        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "Synthetic document 2026")
        pdf.save(pdf_path)
        pdf.close()
        originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (word, excel, pdf_path)}

        window = MainWindow()
        shown = []
        window._show_message = lambda title, text, *args: shown.append((title, text))
        for source, target in ((word, "测试姓名甲"), (excel, "测试姓名甲"), (pdf_path, "2026")):
            window.select_file(source)
            window.mode_input.setCurrentIndex(window.mode_input.findData("find_text"))
            window.request_input.setText(target)
            window.handle_start()
            assert shown[-1][0] == "处理结果" and "找到 1 处" in shown[-1][1], shown[-1]
        window.select_file(excel)
        window.mode_input.setCurrentIndex(window.mode_input.findData("calculate_sum"))
        window.column_input.setText("金额")
        window.handle_start()
        assert "总和是：200" in shown[-1][1], shown[-1]
        assert not settings_path().exists()
        report["checks"].append("no-key GUI local find in DOCX/XLSX/PDF and Excel sum")

        service = SafeModificationService()
        for source, action in ((word, {"action": "replace_text", "target": "测试姓名甲", "replacement": "测试姓名乙"}),
                               (excel, {"action": "replace_cell", "sheet": "报销表", "cell": "B2", "value": 300})):
            plan = validate_operation_plan({"explanation": "合成验收", "operations": [action]})
            preview = service.prepare(source, plan)
            assert not preview.output_path.exists()
            result = service.execute(preview, plan)
            get_processor(result.output_path).read(result.output_path)
            assert hashlib.sha256(result.backup_path.read_bytes()).hexdigest() == originals[source]
            if source == word:
                assert "测试姓名乙" in Document(result.output_path).paragraphs[0].text
            else:
                reopened = load_workbook(result.output_path)
                assert reopened["报销表"]["B2"].value == 300
                reopened.close()
            restored = samples / f"恢复 {source.name}"
            service._backups.restore_to_new_file(result.backup_path, restored)
            assert hashlib.sha256(restored.read_bytes()).hexdigest() == originals[source]
        report["checks"].append("DOCX/XLSX preview, backup, new output, reopen and restore")
        assert all(hashlib.sha256(p.read_bytes()).hexdigest() == h for p, h in originals.items())
        report["original_sha256"] = {p.name: h for p, h in originals.items()}
        report["checks"].append("all three original SHA-256 unchanged")

        test_key = "synthetic-test-key-never-valid"
        DeepSeekSettings(api_key=test_key).save()
        assert test_key not in settings_path().read_text(encoding="utf-8")
        assert DeepSeekSettings.load().api_key == test_key
        DeepSeekSettings(api_key="").save()
        assert DeepSeekSettings.load().api_key == ""
        report["checks"].append("real Windows DPAPI synthetic credential round trip and clear")
        window.show()
        application.processEvents()
        window.grab().save(str(root / "main-window.png"))
        dialog = SettingsDialog(window)
        dialog.show()
        application.processEvents()
        dialog.grab().save(str(root / "settings-dialog.png"))
        dialog.close()
        window.close()
        report["status"] = "passed"
    except Exception as error:
        report["status"] = "failed"
        report["error_type"] = type(error).__name__
        # All inputs to this command are synthetic; no raw configuration is logged.
    (root / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["status"] == "passed" else 1
